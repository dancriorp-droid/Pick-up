"""
PICK-UP SNIPER - App Streamlit com Google Drive
Upload dos HITS → Atualiza planilha (salva no Drive) → Gera PPTXs → Download
"""
import streamlit as st
import os, sys, re, json, shutil, calendar, zipfile, tempfile, io
from datetime import datetime, date, timedelta
from collections import OrderedDict
import pytz

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from lxml import etree

from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload, MediaIoBaseDownload

st.set_page_config(page_title="Pick-up Sniper", page_icon="🎯", layout="wide")

HOTEIS = {
    'normandie': {'aba_planilha': 'Normandie', 'template_pptx': 'pickup_normandie_template.pptx', 'total_uhs': 164},
    'alto': {'aba_planilha': 'Alto da Boa Vista', 'aba_metas': 'Alto da Boa Vista Metas', 'template_pptx': 'pickup_alto_template.pptx', 'total_uhs': 21},
}
MESES_PT = ['Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
FONT_DEFAULT = Font(name='Aptos Narrow', size=11)
FMT_RECEITA = '#,##0.00'
FMT_OCC = '#,##0.00'
METAS = {
    'normandie': {5:{'dm':315.92,'occ':0.85,'rec':1390200},6:{'dm':256.90,'occ':0.82,'rec':1055400},7:{'dm':289.89,'occ':0.72,'rec':1080545},8:{'dm':323.29,'occ':0.66,'rec':1104623.01},9:{'dm':322.94,'occ':0.70,'rec':1132558.24},10:{'dm':307.69,'occ':0.66,'rec':1051334.77},11:{'dm':427.69,'occ':0.78,'rec':1671327.84},12:{'dm':283.21,'occ':0.52,'rec':762401.25},1:{'dm':329.13,'occ':0.41,'rec':698591.23}},
    'alto': {5:{'dm':1239.98,'occ':0.39,'rec':314818.16},6:{'dm':1638.57,'occ':0.38,'rec':392273.26},7:{'dm':1559.03,'occ':0.67,'rec':680000},8:{'dm':1250.82,'occ':0.42,'rec':342000},9:{'dm':1328.23,'occ':0.30,'rec':251034.77},10:{'dm':1168.04,'occ':0.24,'rec':182494.46},11:{'dm':1191.08,'occ':0.32,'rec':240122.57},12:{'dm':1365.42,'occ':0.36,'rec':320000},1:{'dm':1114.39,'occ':0.31,'rec':224895}},
}

# === GOOGLE DRIVE ===
def get_drive_service():
    try:
        creds = Credentials.from_service_account_info(dict(st.secrets["gcp_service_account"]), scopes=['https://www.googleapis.com/auth/drive'])
        return build('drive', 'v3', credentials=creds)
    except:
        return None

def drive_find(svc, name, folder):
    q = f"name='{name}' and trashed=false"
    if folder: q += f" and '{folder}' in parents"
    r = svc.files().list(q=q, fields="files(id,name,modifiedTime)").execute()
    return r.get('files', [None])[0] if r.get('files') else None

def drive_download(svc, fid):
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, svc.files().get_media(fileId=fid))
    done = False
    while not done: _, done = dl.next_chunk()
    buf.seek(0); return buf.read()

def drive_upload(svc, name, data, folder=None, existing_id=None):
    media = MediaIoBaseUpload(io.BytesIO(data), mimetype='application/octet-stream', resumable=False)
    if existing_id:
        svc.files().update(fileId=existing_id, media_body=media).execute()
    else:
        meta = {'name': name}
        if folder: meta['parents'] = [folder]
        svc.files().create(body=meta, media_body=media).execute()

# === PICKUP FUNCTIONS ===
def ler_hits(file_bytes):
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.worksheets[0]; dados = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]: continue
        data_raw=row[2]; total=row[4]; ocupado=row[6]; receita=row[20]
        ocup_kpi=row[31] if len(row)>31 else None
        if not data_raw or not total or ocupado is None: continue
        if isinstance(data_raw, datetime): dt=data_raw.date()
        elif isinstance(data_raw, date): dt=data_raw
        elif isinstance(data_raw, str):
            for fmt in ('%d/%m/%Y','%Y-%m-%d'):
                try: dt=datetime.strptime(data_raw.strip(),fmt).date(); break
                except: continue
            else: continue
        else: continue
        occ=(float(ocupado)/float(total))*100 if float(total)>0 else 0.
        kpi=int(ocup_kpi) if ocup_kpi and isinstance(ocup_kpi,(int,float)) else int(ocupado) if ocupado else 0
        dados[dt]={'receita':float(receita) if receita else 0.,'ocupacao':occ,'ocup_kpi':kpi}
    wb.close(); return dados

def mapear_meses(ws):
    meses=[]
    for c in range(1,ws.max_column+1):
        if ws.cell(row=1,column=c).value=='Dia':
            for r in range(3,ws.max_row+1):
                val=ws.cell(row=r,column=c).value
                if val is None: continue
                if isinstance(val,str) and val.strip().upper()=='TOTAL': break
                try:
                    dt=val if isinstance(val,(datetime,date)) else pd.to_datetime(val)
                    if isinstance(dt,datetime): dt=dt.date()
                    elif isinstance(dt,pd.Timestamp): dt=dt.date()
                    tr=None
                    for t in range(r,ws.max_row+1):
                        tv=ws.cell(row=t,column=c).value
                        if tv and isinstance(tv,str) and tv.strip().upper()=='TOTAL': tr=t; break
                    meses.append({'col_base':c,'mes':dt.month,'ano':dt.year,'total_row':tr}); break
                except: continue
    return meses

def criar_mes_novo(ws,exist,nm,na,hk):
    cn=exist[-1]['col_base']+8 if exist else 1
    nd=calendar.monthrange(na,nm)[1]
    for i,h in enumerate(['Dia','Ontem',None,'Hoje',None,'Pick-up',None,None]):
        ws.cell(row=1,column=cn+i).value=h
    for i,h in enumerate([None,'Receita','Ocupação','Receita','Ocupação','Receita','Ocupação',None]):
        ws.cell(row=2,column=cn+i).value=h
    for d in range(1,nd+1):
        r=2+d; ws.cell(row=r,column=cn).value=datetime(na,nm,d)
        for o in range(1,7): ws.cell(row=r,column=cn+o).value=0
    tr=3+nd; ws.cell(row=tr,column=cn).value='TOTAL'
    m=METAS.get(hk,{}).get(nm,{})
    if m:
        ms=f"Meta: R$ {m['rec']:,.2f}".replace(',','X').replace('.',',').replace('X','.')
        ws.cell(row=tr+3,column=cn+2).value=ms
    return {'col_base':cn,'mes':nm,'ano':na,'total_row':tr}

def atualizar_aba(ws,dados_hits,hk):
    meses=mapear_meses(ws)
    if not meses: return 0,{}
    datas=sorted(dados_hits.keys())
    if not datas: return 0,{}
    mh=set(d.month for d in datas); mp=set(m['mes'] for m in meses); mx=max(mp) if mp else 0
    for mn in sorted(mh-mp):
        if mn<=mx: continue
        an=next(d.year for d in datas if d.month==mn)
        meses.append(criar_mes_novo(ws,meses,mn,an,hk))
    att=0; pu_mes={}
    for mi in meses:
        c=mi['col_base']; pt=0.
        for r in range(3,ws.max_row+1):
            v=ws.cell(row=r,column=c).value
            if v is None: continue
            if isinstance(v,str) and v.strip().upper()=='TOTAL': break
            try:
                if isinstance(v,datetime): dt=v.date()
                elif isinstance(v,date): dt=v
                else: dt=pd.to_datetime(v).date()
            except: continue
            if dt not in dados_hits: continue
            h=dados_hits[dt]
            hr=ws.cell(row=r,column=c+3).value; ho=ws.cell(row=r,column=c+4).value
            or_=float(hr) if hr and isinstance(hr,(int,float)) else 0.
            oo=float(ho) if ho and isinstance(ho,(int,float)) else 0.
            ws.cell(row=r,column=c+1).value=or_; ws.cell(row=r,column=c+2).value=oo
            ws.cell(row=r,column=c+3).value=h['receita']; ws.cell(row=r,column=c+4).value=h['ocupacao']
            pr=h['receita']-or_; po=h['ocupacao']-oo
            ws.cell(row=r,column=c+5).value=pr; ws.cell(row=r,column=c+6).value=po
            pt+=pr; att+=1
        pu_mes[mi['mes']]=pt
    for mi in meses:
        c=mi['col_base']; tr=mi['total_row']
        if not tr: continue
        ss={1:0,3:0,5:0}; aa={2:[],4:[],6:[]}
        for r in range(3,tr):
            v=ws.cell(row=r,column=c).value
            if v is None: continue
            if isinstance(v,str) and v.strip().upper()=='TOTAL': break
            for o in [1,3,5]:
                vv=ws.cell(row=r,column=c+o).value
                if vv and isinstance(vv,(int,float)): ss[o]+=float(vv)
            for o in [2,4,6]:
                vv=ws.cell(row=r,column=c+o).value
                if vv and isinstance(vv,(int,float)): aa[o].append(float(vv))
        for o,t in ss.items(): ws.cell(row=tr,column=c+o).value=t
        for o,v in aa.items(): ws.cell(row=tr,column=c+o).value=sum(v)/len(v) if v else 0.
    for mi in meses:
        c=mi['col_base']; tr=mi['total_row']; trec=tkpi=0
        for r in range(3,ws.max_row+1):
            v=ws.cell(row=r,column=c).value
            if v is None: continue
            if isinstance(v,str) and v.strip().upper()=='TOTAL': break
            try:
                if isinstance(v,datetime): dt=v.date()
                elif isinstance(v,date): dt=v
                else: dt=pd.to_datetime(v).date()
            except: continue
            if dt in dados_hits:
                trec+=dados_hits[dt]['receita']; tkpi+=dados_hits[dt].get('ocup_kpi',0)
        dm=trec/tkpi if tkpi>0 else 0.
        if dm>0 and tr: ws.cell(row=tr+2,column=c+5).value=round(dm,2)
    return att,pu_mes

def atualizar_metas_alto(wb):
    if 'Alto da Boa Vista Metas' not in wb.sheetnames: return
    if 'Alto da Boa Vista' not in wb.sheetnames: return
    wd=wb['Alto da Boa Vista']; wm=wb['Alto da Boa Vista Metas']
    meses=mapear_meses(wd)
    for idx,mi in enumerate(meses):
        if idx>=9: break
        b=7+idx*7; c=mi['col_base']; tr=to=0.; nd=0
        for r in range(3,wd.max_row+1):
            v=wd.cell(row=r,column=c).value
            if v is None: continue
            if isinstance(v,str) and v.strip().upper()=='TOTAL': break
            nd+=1
            rc=wd.cell(row=r,column=c+3).value; oc=wd.cell(row=r,column=c+4).value
            if rc and isinstance(rc,(int,float)): tr+=float(rc)
            if oc and isinstance(oc,(int,float)): to+=float(oc)
        ndm=calendar.monthrange(mi['ano'],mi['mes'])[1]
        om=to/ndm if ndm>0 else 0.
        wm.cell(row=b+3,column=12).value=tr
        wm.cell(row=b+4,column=12).value=om/100 if om>1 else om
        mr=wm.cell(row=b+3,column=13).value; mo=wm.cell(row=b+4,column=13).value
        if mr and isinstance(mr,(int,float)) and float(mr)>0: wm.cell(row=b+3,column=14).value=tr/float(mr)
        if mo and isinstance(mo,(int,float)) and float(mo)>0: wm.cell(row=b+4,column=14).value=(om/100 if om>1 else om)/float(mo)
        dv=wm.cell(row=b+5,column=12).value; md=wm.cell(row=b+5,column=13).value
        if dv and md and isinstance(md,(int,float)) and float(md)>0: wm.cell(row=b+5,column=14).value=float(dv)/float(md)
        br_tz=pytz.timezone('America/Sao_Paulo'); hj=datetime.now(br_tz).date()
        ud=date(mi['ano'],mi['mes'],calendar.monthrange(mi['ano'],mi['mes'])[1])
        df=max(1,(ud-hj).days+1); wm.cell(row=b,column=13).value=df
        if mr and isinstance(mr,(int,float)) and float(mr)>0: wm.cell(row=b+3,column=15).value=(float(mr)-tr)/df

# === INTERFACE ===
st.title("🎯 PICK-UP SNIPER")
drive_svc = get_drive_service()
FOLDER_ID = st.secrets.get("drive_folder_id","") if "drive_folder_id" in st.secrets else ""
drive_ok = drive_svc is not None and bool(FOLDER_ID)

if drive_ok: st.success("☁️ Google Drive conectado!")
else: st.info("📁 Modo manual — envie a planilha por upload")

st.markdown("---")
c1,c2=st.columns(2)
with c1: hits_n=st.file_uploader("🏩 HITS Normandie",type=['xlsx'],key="hn")
with c2: hits_a=st.file_uploader("🏨 HITS Alto",type=['xlsx'],key="ha")

plan_f=None
if not drive_ok: plan_f=st.file_uploader("📊 Planilha Pick-up_sniper.xlsx",type=['xlsx'],key="pl")

with st.expander("📄 Templates + Histórico (opcional)"):
    t1,t2=st.columns(2)
    with t1: tpl_n=st.file_uploader("Template Normandie",type=['pptx'],key="t1")
    with t2: tpl_a=st.file_uploader("Template Alto",type=['pptx'],key="t2")
    hist_f=st.file_uploader("pickup_historico.json",type=['json'],key="hf")

st.markdown("---")

if st.button("🚀 ATUALIZAR PICK-UP",type="primary",use_container_width=True):
    if not hits_n and not hits_a: st.error("⚠️ Envie pelo menos um HITS!")
    elif not drive_ok and not plan_f: st.error("⚠️ Envie a planilha!")
    else:
        with st.spinner("⏳ Processando..."):
            tmp=tempfile.mkdtemp()
            if drive_ok:
                ff=drive_find(drive_svc,'Pick-up_sniper.xlsx',FOLDER_ID)
                if ff: pd_=drive_download(drive_svc,ff['id']); st.info("☁️ Planilha do Drive")
                else: st.error("❌ Planilha não encontrada no Drive!"); st.stop()
            else: pd_=plan_f.getvalue()
            pp=os.path.join(tmp,'Pick-up_sniper.xlsx')
            with open(pp,'wb') as f: f.write(pd_)

            hist={}
            if drive_ok:
                hff=drive_find(drive_svc,'pickup_historico.json',FOLDER_ID)
                if hff:
                    try: hist=json.loads(drive_download(drive_svc,hff['id']))
                    except: pass
            if hist_f:
                try: hist=json.loads(hist_f.getvalue())
                except: pass

            br=pytz.timezone('America/Sao_Paulo'); hj=datetime.now(br).date(); dw=hj.weekday()
            if dw==0:
                s,sa,d=hj-timedelta(days=3),hj-timedelta(days=2),hj-timedelta(days=1)
                de=f"{s.day}, {sa.day}, {d.strftime('%d/%m')}"
            else: de=(hj-timedelta(days=1)).strftime('%d/%m')

            wb=load_workbook(pp); res={}
            for hk,hf in [('normandie',hits_n),('alto',hits_a)]:
                if not hf: continue
                cfg=HOTEIS[hk]
                if cfg['aba_planilha'] not in wb.sheetnames: continue
                dados=ler_hits(hf.getvalue()); ws=wb[cfg['aba_planilha']]
                n,pm=atualizar_aba(ws,dados,hk)
                if hk not in hist: hist[hk]={}
                for m,pu in pm.items():
                    mk=str(m)
                    if mk not in hist[hk]: hist[hk][mk]=[]
                    if not hist[hk][mk] or hist[hk][mk][-1]['data']!=de:
                        hist[hk][mk].append({'data':de,'pickup_total':round(pu,2)})
                    hist[hk][mk]=hist[hk][mk][-30:]
                res[hk]={'n':n,'pu':pm}

            if 'alto' in res: atualizar_metas_alto(wb)
            wb.save(pp)

            hp=os.path.join(tmp,'pickup_historico.json')
            with open(hp,'w') as f: json.dump(hist,f,indent=2,ensure_ascii=False)

            gs=os.path.join(os.path.dirname(__file__),'gerar_pickup_pptx.py'); pptxs={}
            for hk in res:
                cfg=HOTEIS[hk]; td=None
                tu=tpl_n if hk=='normandie' else tpl_a
                if tu: td=tu.getvalue()
                elif drive_ok:
                    tf=drive_find(drive_svc,cfg['template_pptx'],FOLDER_ID)
                    if tf: td=drive_download(drive_svc,tf['id'])
                if not td and os.path.exists(cfg['template_pptx']):
                    with open(cfg['template_pptx'],'rb') as f: td=f.read()
                if td:
                    tp=os.path.join(tmp,cfg['template_pptx'])
                    with open(tp,'wb') as f: f.write(td)
                    op=os.path.join(tmp,f'pickup_{hk}_gerado.pptx')
                    if os.path.exists(gs):
                        import subprocess
                        subprocess.run([sys.executable,gs,pp,tp,hk,op,hp],capture_output=True,timeout=120)
                    if os.path.exists(op): pptxs[hk]=op

            if drive_ok:
                with open(pp,'rb') as f: pb=f.read()
                ex=drive_find(drive_svc,'Pick-up_sniper.xlsx',FOLDER_ID)
                drive_upload(drive_svc,'Pick-up_sniper.xlsx',pb,FOLDER_ID,ex['id'] if ex else None)
                hb=json.dumps(hist,indent=2,ensure_ascii=False).encode('utf-8')
                eh=drive_find(drive_svc,'pickup_historico.json',FOLDER_ID)
                drive_upload(drive_svc,'pickup_historico.json',hb,FOLDER_ID,eh['id'] if eh else None)
                for hk,path in pptxs.items():
                    with open(path,'rb') as f: xb=f.read()
                    fn=f'pickup_{hk}_gerado.pptx'
                    ep=drive_find(drive_svc,fn,FOLDER_ID)
                    drive_upload(drive_svc,fn,xb,FOLDER_ID,ep['id'] if ep else None)
                st.success("☁️ Tudo salvo no Google Drive!")

        st.success("✅ Concluído!")
        for hk,i in res.items(): st.markdown(f"**{hk.upper()}**: {i['n']} dias atualizados")

        st.markdown("### 📥 Downloads")
        d1,d2=st.columns(2)
        with open(pp,'rb') as f: d1.download_button("📊 Planilha",f.read(),"Pick-up_sniper.xlsx")
        d1.download_button("📋 Histórico",json.dumps(hist,indent=2,ensure_ascii=False),"pickup_historico.json")
        for hk,path in pptxs.items():
            with open(path,'rb') as f:
                nm="Normandie" if hk=='normandie' else "Alto"
                d2.download_button(f"📥 PPTX {nm}",f.read(),f"pickup_{hk}_gerado.pptx")
